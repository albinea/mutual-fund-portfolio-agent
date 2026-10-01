"""Safe, template-based parsing for portfolio CSV and XLSX statement uploads."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from .models import (
    FundDisclosureHolding,
    FundDisclosureImport,
    ImportedHolding,
    PortfolioImport,
)


MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_HOLDINGS = 250
MAX_DISCLOSURE_HOLDINGS = 5_000
REQUIRED_COLUMNS = {"fund_name", "invested_amount", "current_value"}
COLUMN_ALIASES = {
    "fund_name": {"fund", "fund_name", "scheme", "scheme_name", "mutual_fund"},
    "category": {"category", "fund_category", "scheme_category"},
    "units": {"units", "unit", "quantity"},
    "invested_amount": {"invested", "invested_amount", "investment", "purchase_value", "cost_value"},
    "current_value": {"current_value", "current", "market_value", "value", "present_value"},
}
DISCLOSURE_REQUIRED_COLUMNS = {"fund_name", "company_name", "holding_weight"}
DISCLOSURE_COLUMN_ALIASES = {
    "fund_name": {"fund", "fund_name", "scheme", "scheme_name", "mutual_fund"},
    "company_name": {
        "company",
        "company_name",
        "security",
        "security_name",
        "issuer",
        "issuer_name",
        "stock_name",
    },
    "isin": {"isin", "security_isin"},
    "sector": {"sector", "industry", "sector_name"},
    "holding_weight": {
        "holding_weight",
        "weight",
        "weight_percentage",
        "allocation_percent",
        "allocation_percentage",
        "portfolio_weight",
        "portfolio_weight_percentage",
        "percentage",
    },
}


class StatementImportError(ValueError):
    """A client-safe error explaining why a statement cannot be imported."""


class DisclosureImportError(ValueError):
    """A client-safe error explaining why a fund disclosure cannot be imported."""


def _normalized_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")


def normalize_fund_name(value: str) -> str:
    """Create a conservative, case-insensitive key for matching uploaded fund names."""

    return _normalized_header(value)


def _column_map(headers: list[object]) -> dict[str, int]:
    normalized = {_normalized_header(header): index for index, header in enumerate(headers)}
    mapping: dict[str, int] = {}
    for target, aliases in COLUMN_ALIASES.items():
        match = next((alias for alias in aliases if alias in normalized), None)
        if match is not None:
            mapping[target] = normalized[match]
    missing = REQUIRED_COLUMNS - mapping.keys()
    if missing:
        friendly = ", ".join(sorted(missing))
        raise StatementImportError(
            f"The file must contain these columns: fund_name, invested_amount, and current_value. Missing: {friendly}."
        )
    return mapping


def _as_decimal(value: object, field: str, row_number: int, *, default: Decimal | None = None) -> Decimal:
    if value in (None, ""):
        if default is not None:
            return default
        raise StatementImportError(f"Row {row_number}: {field} is required.")
    text = str(value).strip().replace(",", "").replace("₹", "").replace("%", "")
    text = re.sub(r"\bINR\b", "", text, flags=re.IGNORECASE).strip()
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError):
        raise StatementImportError(f"Row {row_number}: {field} must be a valid number.")
    if not amount.is_finite() or amount < 0:
        raise StatementImportError(f"Row {row_number}: {field} cannot be negative.")
    return amount


def _parse_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    non_empty_rows = [row for row in rows if any(value not in (None, "") for value in row)]
    if len(non_empty_rows) < 2:
        raise StatementImportError("The statement must have a header row and at least one holding.")
    mapping = _column_map(non_empty_rows[0])
    parsed: list[dict[str, object]] = []
    for row_number, row in enumerate(non_empty_rows[1:], start=2):
        def value(column: str, default: object = "") -> object:
            index = mapping.get(column)
            return row[index] if index is not None and index < len(row) else default

        fund_name = str(value("fund_name")).strip()
        if not fund_name:
            raise StatementImportError(f"Row {row_number}: fund_name is required.")
        if len(fund_name) > 255:
            raise StatementImportError(f"Row {row_number}: fund_name is too long.")
        category = str(value("category", "Imported")).strip() or "Imported"
        parsed.append(
            {
                "fund_name": fund_name,
                "category": category[:100],
                "units": _as_decimal(value("units", 0), "units", row_number, default=Decimal("0")),
                "invested_amount": _as_decimal(value("invested_amount"), "invested_amount", row_number),
                "current_value": _as_decimal(value("current_value"), "current_value", row_number),
            }
        )
    if not parsed:
        raise StatementImportError("No holdings were found in the uploaded statement.")
    if len(parsed) > MAX_HOLDINGS:
        raise StatementImportError(f"A statement can contain at most {MAX_HOLDINGS} holdings.")
    return parsed


def _rows_from_csv(content: bytes) -> list[list[object]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise StatementImportError("CSV files must be UTF-8 encoded.")
    return list(csv.reader(io.StringIO(decoded)))


def _rows_from_xlsx(content: bytes) -> list[list[object]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise StatementImportError("The XLSX workbook expands beyond the safe 10 MB limit.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    except StatementImportError:
        raise
    except Exception as error:
        raise StatementImportError("The XLSX file could not be read. Use a standard, unprotected Excel workbook.") from error


def parse_statement(uploaded_file) -> tuple[str, list[dict[str, object]]]:
    """Parse one small, known-format statement without storing the original upload."""

    suffix = Path(uploaded_file.name).suffix.casefold()
    if suffix not in {".csv", ".xlsx"}:
        raise StatementImportError("Only CSV and XLSX statements are supported. Export a PDF statement as CSV or XLSX first.")
    if not uploaded_file.size:
        raise StatementImportError("The uploaded file is empty.")
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise StatementImportError("The uploaded file is too large. The limit is 2 MB.")
    content = uploaded_file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise StatementImportError("The uploaded file is too large. The limit is 2 MB.")
    rows = _rows_from_csv(content) if suffix == ".csv" else _rows_from_xlsx(content)
    return suffix.removeprefix("."), _parse_rows(rows)


def _disclosure_column_map(headers: list[object]) -> dict[str, int]:
    normalized = {_normalized_header(header): index for index, header in enumerate(headers)}
    mapping: dict[str, int] = {}
    for target, aliases in DISCLOSURE_COLUMN_ALIASES.items():
        match = next((alias for alias in aliases if alias in normalized), None)
        if match is not None:
            mapping[target] = normalized[match]
    missing = DISCLOSURE_REQUIRED_COLUMNS - mapping.keys()
    if missing:
        friendly = ", ".join(sorted(missing))
        raise DisclosureImportError(
            "The disclosure file must contain these columns: fund_name, company_name, and holding_weight. "
            f"Missing: {friendly}."
        )
    return mapping


def _parse_disclosure_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    non_empty_rows = [row for row in rows if any(value not in (None, "") for value in row)]
    if len(non_empty_rows) < 2:
        raise DisclosureImportError("The disclosure file must have a header row and at least one company holding.")
    mapping = _disclosure_column_map(non_empty_rows[0])
    parsed: list[dict[str, object]] = []
    for row_number, row in enumerate(non_empty_rows[1:], start=2):
        def value(column: str, default: object = "") -> object:
            index = mapping.get(column)
            return row[index] if index is not None and index < len(row) else default

        fund_name = str(value("fund_name")).strip()
        company_name = str(value("company_name")).strip()
        if not fund_name:
            raise DisclosureImportError(f"Row {row_number}: fund_name is required.")
        if not company_name:
            raise DisclosureImportError(f"Row {row_number}: company_name is required.")
        if len(fund_name) > 255 or len(company_name) > 255:
            raise DisclosureImportError(f"Row {row_number}: fund_name and company_name must be 255 characters or fewer.")
        try:
            holding_weight = _as_decimal(value("holding_weight"), "holding_weight", row_number)
        except StatementImportError as error:
            raise DisclosureImportError(str(error)) from error
        if holding_weight > 100:
            raise DisclosureImportError(f"Row {row_number}: holding_weight cannot exceed 100%.")
        isin = re.sub(r"\s+", "", str(value("isin")).upper())
        sector = str(value("sector")).strip()
        parsed.append(
            {
                "fund_name": fund_name,
                "normalized_fund_name": normalize_fund_name(fund_name),
                "company_name": company_name,
                "isin": isin[:32],
                "sector": sector[:120],
                "holding_weight": holding_weight,
            }
        )
    if len(parsed) > MAX_DISCLOSURE_HOLDINGS:
        raise DisclosureImportError(f"A disclosure upload can contain at most {MAX_DISCLOSURE_HOLDINGS} company holdings.")
    return parsed


def parse_fund_disclosures(uploaded_file) -> tuple[str, list[dict[str, object]]]:
    """Parse a normalized CSV/XLSX disclosure file without retaining the original upload."""

    suffix = Path(uploaded_file.name).suffix.casefold()
    if suffix not in {".csv", ".xlsx"}:
        raise DisclosureImportError("Only CSV and XLSX disclosure files are supported.")
    if not uploaded_file.size:
        raise DisclosureImportError("The uploaded disclosure file is empty.")
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise DisclosureImportError("The uploaded disclosure file is too large. The limit is 2 MB.")
    content = uploaded_file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise DisclosureImportError("The uploaded disclosure file is too large. The limit is 2 MB.")
    try:
        rows = _rows_from_csv(content) if suffix == ".csv" else _rows_from_xlsx(content)
    except StatementImportError as error:
        raise DisclosureImportError(str(error)) from error
    return suffix.removeprefix("."), _parse_disclosure_rows(rows)


def _combine_duplicate_holdings(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    combined: dict[tuple[str, str], dict[str, object]] = defaultdict(
        lambda: {"units": Decimal("0"), "invested_amount": Decimal("0"), "current_value": Decimal("0")}
    )
    for row in rows:
        key = (str(row["fund_name"]), str(row["category"]))
        item = combined[key]
        item["fund_name"] = key[0]
        item["category"] = key[1]
        item["units"] += row["units"]
        item["invested_amount"] += row["invested_amount"]
        item["current_value"] += row["current_value"]
    return list(combined.values())


def create_portfolio_import(*, user_id: str, uploaded_file, snapshot_date: date | None = None) -> PortfolioImport:
    source_format, rows = parse_statement(uploaded_file)
    holdings = _combine_duplicate_holdings(rows)
    with transaction.atomic():
        portfolio_import = PortfolioImport.objects.create(
            user_id=user_id,
            source_file_name=Path(uploaded_file.name).name[:255],
            source_format=source_format,
            snapshot_date=snapshot_date or date.today(),
            is_active=False,
        )
        ImportedHolding.objects.bulk_create(
            [ImportedHolding(portfolio_import=portfolio_import, **holding) for holding in holdings]
        )
        PortfolioImport.objects.filter(user_id=user_id, is_active=True).update(is_active=False)
        active_import = (
            PortfolioImport.objects.filter(user_id=user_id)
            .order_by("-snapshot_date", "-id")
            .first()
        )
        active_import.is_active = True
        active_import.save(update_fields=["is_active"])
        portfolio_import.refresh_from_db(fields=["is_active"])
    return portfolio_import


def _combine_duplicate_disclosures(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    combined: dict[tuple[str, str], dict[str, object]] = {}
    for row in rows:
        company_key = str(row["isin"]) or _normalized_header(row["company_name"])
        key = (str(row["normalized_fund_name"]), company_key)
        if key not in combined:
            combined[key] = row.copy()
            continue
        item = combined[key]
        item["holding_weight"] += row["holding_weight"]
        if not item["sector"] and row["sector"]:
            item["sector"] = row["sector"]
    return list(combined.values())


def create_fund_disclosure_import(
    *,
    uploaded_file,
    disclosure_date: date | None = None,
    source_name: str | None = None,
    source_url: str | None = None,
) -> FundDisclosureImport:
    """Store a manually curated disclosure dataset for use in portfolio overlap calculations."""

    source_format, rows = parse_fund_disclosures(uploaded_file)
    holdings = _combine_duplicate_disclosures(rows)
    with transaction.atomic():
        disclosure_import = FundDisclosureImport.objects.create(
            source_file_name=Path(uploaded_file.name).name[:255],
            source_format=source_format,
            source_name=(source_name or "Manual disclosure upload")[:255],
            source_url=source_url or "",
            disclosure_date=disclosure_date or date.today(),
        )
        FundDisclosureHolding.objects.bulk_create(
            [FundDisclosureHolding(disclosure_import=disclosure_import, **holding) for holding in holdings]
        )
    return disclosure_import
